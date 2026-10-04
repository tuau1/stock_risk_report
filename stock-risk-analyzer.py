#----------- IMPORTS ------------

import pandas as pd
import numpy as np
import yfinance as yf
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image

#----------- SETTINGS ---------

TICKER = "SPY"
FOLDER = ""
PERIOD = "3y"

#-----------DATA----------

df = yf.download(TICKER, period=PERIOD, auto_adjust=True)

df.columns = df.columns.droplevel(1)

df = df.dropna(subset=["Close"])

#----------COLUMNS--------

df["daily_return"] = df["Close"].pct_change()

annualized_volatility = df["daily_return"].std() * np.sqrt(252)

average_daily_return = df["daily_return"].mean()

first = df["Close"].iloc[0]

last = df["Close"].iloc[-1]

best_day = df["daily_return"].max()

worst_day = df["daily_return"].min()

total = ( last - first ) / first

up_days = len(df[df["daily_return"] > 0])

up_days_pct = up_days / len(df)

daily_volatility = df["daily_return"].std()




#--------SUMMARY------------

summary = pd.DataFrame({
    "Metric": [
       
        "Annualized Volatility",
        "Average Daily Return",
        "Best Day",
        "Worst Day",
        "Total Return",
        "Up Days",
        "Up Days Pct"],


    "Value": [
       
        annualized_volatility,
        average_daily_return,
        best_day,
        worst_day,
        total,
        up_days,
        up_days_pct]})



#----------PLOT------------
plt.figure(figsize=(10, 5))

plt.plot(df["daily_return"])
plt.title(TICKER + " Daily Returns - " + PERIOD)

plt.axhline(0, color="black", linewidth=0.8)
plt.axhline(daily_volatility, color="red", linewidth= 1, linestyle="--")
plt.axhline(-daily_volatility, color="red", linewidth= 1, linestyle="--")


plt.xlabel("Date")

plt.ylabel("Daily Return")

plt.grid(True)

plt.xticks(rotation=45)


plt.savefig(FOLDER + TICKER + "_returns.png", dpi=150, bbox_inches="tight")



with pd.ExcelWriter(FOLDER + TICKER + "_report.xlsx") as writer:
    df.to_excel(writer, sheet_name="Data")
    summary.to_excel(writer,sheet_name="Summary",index=False)
    

wb = load_workbook(FOLDER + TICKER + "_report.xlsx")

ws = wb.create_sheet("Chart")

ws.add_image(Image(FOLDER + TICKER + "_returns.png"), "B2")  

wb.save(FOLDER + TICKER + "_report.xlsx")

#----------FILE-------------




print("created")


























                 
